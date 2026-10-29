"""Bounded local media retrieval, document reading, and offline speech transcription."""
from __future__ import annotations
import argparse
import contextlib
import hashlib
import io
import json
from pathlib import Path
import shutil
import sqlite3
import subprocess
import tempfile
import threading
import wave
import zipfile
import xml.etree.ElementTree as ET

from wechat_index import ROOT, DATA, INDEX, connect, resolve_session, page, now
from session_policy import excluded_ids, exclude_public_accounts, is_public_account
from media_sources import current_account, ImageReader
from mcp.server.mcpserver.exceptions import ToolError

MEDIA_INDEX = DATA / 'wechat-media-index.sqlite'
OUTPUT = ROOT / 'exports/media'
MODEL = ROOT / 'models/faster-whisper-base'
_ASR = None
_ASR_LOCK = threading.Lock()


def _active(identity):
    with contextlib.closing(connect()) as c:
        meta = {r['key']:json.loads(r['value']) for r in c.execute('SELECT * FROM metadata')}
        r = c.execute('''SELECT m.*,s.display AS session_display FROM messages m
            JOIN sessions s ON s.username=m.session_id WHERE m.identity=?''',(identity,)).fetchone()
    if not r:
        raise ToolError('消息不存在于当前有效索引，或所属会话已被排除。')
    r = dict(r)
    if r['session_id'] in excluded_ids(meta['account_fingerprint']) or (
        exclude_public_accounts(meta['account_fingerprint']) and is_public_account(r['session_id'])):
        raise ToolError('该会话已被排除。')
    return r, meta


def _db(write=False):
    if not MEDIA_INDEX.exists():
        raise ValueError('请先更新媒体索引：build_media_index.py。')
    c = sqlite3.connect(MEDIA_INDEX if write else MEDIA_INDEX.as_uri()+'?mode=ro',uri=not write)
    c.row_factory = sqlite3.Row
    return c


def media_status():
    if not MEDIA_INDEX.exists():
        return dict(ready=False)
    with contextlib.closing(_db()) as c, contextlib.closing(connect()) as active:
        meta = json.loads(c.execute("SELECT value FROM metadata WHERE key='status'").fetchone()[0])
        current = {r['key']:json.loads(r['value']) for r in active.execute('SELECT * FROM metadata')}
        if meta['account_fingerprint'] != current['account_fingerprint']:
            raise ValueError('媒体索引账号与当前索引不一致。')
        c.execute('ATTACH DATABASE ? AS current_index',(INDEX.as_uri()+'?mode=ro',))
        counts = [dict(r) for r in c.execute('''SELECT a.kind,json_extract(a.details,'$.status') AS status,COUNT(*) AS count
            FROM assets a JOIN current_index.messages m ON m.identity=a.identity GROUP BY a.kind,status''')]
    return dict(ready=True,**meta,current_message_indexed_at=current['indexed_at'],
                needs_refresh=meta['message_indexed_at']!=current['indexed_at'],current_counts=counts,
                local_transcription_model_ready=(MODEL/'model.bin').exists(),
                capabilities=['image_decode','silk_to_wav','local_speech_to_text','video_frames','video_audio_transcription','document_text'],
                output_dir=str(OUTPUT))


def list_media(session='',kind='',availability='',limit=50,offset=0):
    page(limit,offset)
    if kind and kind not in ('image','voice','video','file'):
        raise ValueError('kind应为image、voice、video或file。')
    if availability and availability not in ('available','missing','candidate'):
        raise ValueError('availability应为available、missing或candidate。')
    with contextlib.closing(connect()) as c, contextlib.closing(_db()) as media:
        current_fp = json.loads(c.execute("SELECT value FROM metadata WHERE key='account_fingerprint'").fetchone()[0])
        info = json.loads(media.execute("SELECT value FROM metadata WHERE key='status'").fetchone()[0])
        if info['account_fingerprint'] != current_fp:
            raise ValueError('媒体索引账号不一致。')
        c.execute('ATTACH DATABASE ? AS media_index',(MEDIA_INDEX.as_uri()+'?mode=ro',))
        clause=[];args=[]
        if session:
            clause.append('m.session_id=?');args.append(resolve_session(c,session)['username'])
        if kind:
            clause.append('m.kind=?');args.append(kind)
        if availability:
            clause.append("json_extract(a.details,'$.status')=?");args.append(availability)
        excluded=excluded_ids(current_fp)
        if excluded:
            clause.append('m.session_id NOT IN ('+','.join('?' for _ in excluded)+')');args.extend(sorted(excluded))
        if exclude_public_accounts(current_fp):
            clause.append("substr(m.session_id,1,3)<>'gh_'")
        base=''' FROM media_index.assets a JOIN messages m ON m.identity=a.identity
            JOIN sessions s ON s.username=m.session_id'''+(' WHERE '+' AND '.join(clause) if clause else '')
        total=c.execute('SELECT COUNT(*)'+base,args).fetchone()[0]
        rows=c.execute('''SELECT m.identity,m.time_str,m.kind,m.is_sender,s.display AS session_display,a.details'''+base+
                       ' ORDER BY m.create_time DESC,m.identity LIMIT ? OFFSET ?',args+[limit,offset]).fetchall()
    result=[]
    for r in rows:
        r=dict(r);detail=json.loads(r.pop('details'))
        r.update(status=detail['status'],reason=detail['reason'],
                 tiers=sorted({x['tier'] for x in detail.get('candidates',[])}))
        result.append(r)
    return dict(total=total,returned_count=len(result),offset=offset,messages=result)


def _transcribe(path,language='zh'):
    global _ASR
    if not (MODEL/'model.bin').exists():
        raise ValueError('本地语音转写模型尚未准备好。')
    from faster_whisper import WhisperModel
    with _ASR_LOCK:
        if _ASR is None:
            _ASR=WhisperModel(str(MODEL),device='cpu',compute_type='int8',cpu_threads=4,local_files_only=True)
        segments,info=_ASR.transcribe(str(path),language=language or None,beam_size=5,
            condition_on_previous_text=False,vad_filter=True)
        segments=[dict(start=round(s.start,2),end=round(s.end,2),text=s.text) for s in segments]
    return dict(text=''.join(s['text'] for s in segments),segments=segments,
                language=info.language,model='faster-whisper-base/int8',automatic_transcription=True)


def _voice(blob,folder,transcribe,language):
    import pilk
    silk=folder/'audio.silk';wav=folder/'audio.wav'
    silk.write_bytes(blob)
    with tempfile.TemporaryDirectory(prefix='voice-',dir=DATA) as tmp:
        pcm=Path(tmp)/'voice.pcm'
        pilk.decode(str(silk),str(pcm),pcm_rate=24000)
        raw=pcm.read_bytes()
        if not raw:
            raise ValueError('语音解码为空。')
        with wave.open(str(wav),'wb') as w:
            w.setnchannels(1);w.setsampwidth(2);w.setframerate(24000);w.writeframes(raw)
    result=dict(status='ready',path=str(wav),silk_path=str(silk),duration_seconds=len(raw)/48000)
    if transcribe:
        result['transcription']=_transcribe(wav,language)
    return result


def _video(path,folder,frames,transcribe,language):
    import av
    with av.open(str(path)) as c:
        stream=c.streams.video[0]
        duration=float(stream.duration*stream.time_base) if stream.duration else float(c.duration or 0)/av.time_base
        has_audio=bool(c.streams.audio)
        width,height=stream.width,stream.height
        outputs=[]
        for i in range(frames):
            target=duration*(i+1)/(frames+1) if duration else 0
            c.seek(max(0,int(target/float(stream.time_base))),stream=stream,backward=True)
            for f in c.decode(stream):
                if f.time is None or f.time>=target:
                    p=folder/f'frame-{i+1:02d}.jpg'
                    image=f.to_image();image.thumbnail((1600,1600));image.save(p,quality=90)
                    outputs.append(dict(path=str(p),time_seconds=round(float(f.time or target),2)))
                    break
    result=dict(status='ready',path=str(path),duration_seconds=duration,width=width,height=height,
                frames=outputs,has_audio=has_audio)
    if transcribe and has_audio:
        if duration>600:
            result['transcription_status']='超过10分钟，未自动转写'
        else:
            result['transcription']=_transcribe(path,language)
    return result


def _document(path,max_chars):
    ext=path.suffix.lower();parts=[]
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(str(path))
        for i,p in enumerate(reader.pages[:20],1):
            parts.append(f'--- 第{i}页 ---\n'+(p.extract_text() or ''))
            if sum(map(len,parts))>=max_chars:break
    elif ext=='.docx':
        from docx import Document
        d=Document(str(path))
        for p in d.paragraphs:
            parts.append(p.text)
            if sum(map(len,parts))>=max_chars:break
        if sum(map(len,parts))<max_chars:
            for t in d.tables:
                for row in t.rows:
                    parts.append('\t'.join(c.text for c in row.cells))
                    if sum(map(len,parts))>=max_chars:break
                if sum(map(len,parts))>=max_chars:break
    elif ext in ('.xlsx','.xlsm'):
        from openpyxl import load_workbook
        book=load_workbook(str(path),read_only=True,data_only=True)
        try:
            for sheet in book.worksheets[:5]:
                parts.append('--- '+sheet.title+' ---')
                for row in sheet.iter_rows(max_row=1000,max_col=30,values_only=True):
                    parts.append('\t'.join('' if v is None else str(v) for v in row))
                    if sum(map(len,parts))>=max_chars:break
                if sum(map(len,parts))>=max_chars:break
        finally:book.close()
    elif ext=='.xls':
        import xlrd
        book=xlrd.open_workbook(str(path),on_demand=True)
        try:
            for sheet in book.sheets()[:5]:
                parts.append('--- '+sheet.name+' ---')
                for i in range(min(sheet.nrows,1000)):
                    parts.append('\t'.join(str(v) for v in sheet.row_values(i)[:30]))
                    if sum(map(len,parts))>=max_chars:break
                if sum(map(len,parts))>=max_chars:break
        finally:book.release_resources()
    elif ext in ('.doc','.ppt'):
        soffice=next((p for p in [Path('C:/Program Files/LibreOffice/program/soffice.exe'),
            Path('D:/Program Files/LibreOffice/program/soffice.exe')] if p.exists()),None)
        if soffice is None:
            return dict(status='converter_unavailable',extension=ext,text='')
        with tempfile.TemporaryDirectory(prefix='document-',dir=DATA) as tmp:
            temp=Path(tmp);converted=temp/'converted';converted.mkdir()
            target_ext='docx' if ext=='.doc' else 'pptx'
            run=subprocess.run([str(soffice),'-env:UserInstallation='+(temp/'profile').as_uri(),
                '--headless','--convert-to',target_ext,'--outdir',str(converted),str(path)],
                capture_output=True,timeout=90,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            candidates=list(converted.glob('*.'+target_ext))
            if run.returncode or len(candidates)!=1:
                return dict(status='conversion_failed',extension=ext,text='')
            result=_document(candidates[0],max_chars)
            result['converted_from']=ext
            return result
    elif ext=='.pptx':
        with zipfile.ZipFile(path) as z:
            import re
            names=sorted((n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)),
                         key=lambda n:int(re.search(r'slide(\d+)',n).group(1)))
            for n in names[:30]:
                if z.getinfo(n).file_size>10*1024*1024:continue
                root=ET.fromstring(z.read(n))
                parts.append('\n'.join(x.text or '' for x in root.iter() if x.tag.endswith('}t')))
                if sum(map(len,parts))>=max_chars:break
    elif ext in ('.txt','.md','.csv','.tsv','.json','.log','.xml','.html','.htm','.py','.js','.srt'):
        with path.open('rb') as f:raw=f.read(max_chars*4+4)
        for enc in ('utf-8-sig','gb18030','utf-16'):
            try:parts=[raw.decode(enc)];break
            except UnicodeError:continue
        if not parts:parts=[raw.decode('utf-8','replace')]
    else:
        return dict(status='unsupported_format',extension=ext,text='')
    text='\n'.join(parts)
    return dict(status='ready' if text.strip() else 'no_extractable_text',text=text[:max_chars],
                truncated=len(text)>max_chars,extraction_limits='PDF前20页；表格前5个工作表、每表前1000行30列；PPT前30页；按字符上限截断')


def _copy_local(source,target):
    # WeChat caches can have Windows read-only attributes. Do not copy those
    # attributes onto outputs, and reuse an identical existing export.
    if target.exists():
        with source.open('rb') as f:digest=hashlib.file_digest(f,'sha256').hexdigest()
        with target.open('rb') as f:existing=hashlib.file_digest(f,'sha256').hexdigest()
        if digest==existing:
            return target
        target=target.with_name(target.stem+'-'+digest[:12]+target.suffix)
        if target.exists():
            with target.open('rb') as f:
                if hashlib.file_digest(f,'sha256').hexdigest()==digest:return target
            raise ToolError('已有导出内容被修改，保留文件并停止覆盖。')
    shutil.copyfile(source,target)
    return target


def read_media(identity,transcribe=True,language='zh',frames=3,max_chars=20000):
    if not 1<=frames<=6 or not 1000<=max_chars<=50000:
        raise ValueError('frames应为1—6；max_chars应为1000—50000。')
    message,meta=_active(identity)
    account=current_account()
    with contextlib.closing(_db()) as c:
        info=json.loads(c.execute("SELECT value FROM metadata WHERE key='status'").fetchone()[0])
        if info['account_fingerprint']!=meta['account_fingerprint']:
            raise ValueError('媒体索引账号不一致。')
        row=c.execute('SELECT * FROM assets WHERE identity=?',(identity,)).fetchone()
        cached=c.execute('SELECT result FROM results WHERE identity=?',(identity,)).fetchone()
    if row is None:
        return dict(status='not_indexed',reason='尚未建立此消息的媒体映射，请更新媒体索引。',identity=identity)
    detail=json.loads(row['details'])
    if detail['status']!='available':
        return dict(identity=identity,kind=row['kind'],**detail)
    signature=hashlib.sha256((row['kind']+row['details']).encode()+(row['voice'] or b'')).hexdigest()
    options=dict(transcribe=transcribe,language=language,frames=frames,max_chars=max_chars,reader_version=2)
    if cached:
        previous=json.loads(cached['result'])
        if previous.get('asset_signature')==signature and previous.get('options')==options and Path(previous.get('path','')).is_file():
            return {**previous,'cached':True}
    folder=OUTPUT/hashlib.sha256(identity.encode()).hexdigest()[:24]
    folder.mkdir(parents=True,exist_ok=True)
    base=dict(identity=identity,kind=row['kind'],session=message['session_display'],time=message['time_str'],processed_at=now())
    if row['kind']=='voice':
        result=_voice(row['voice'],folder,transcribe,language)
        provenance=dict(match=detail['reason'],sha256=hashlib.sha256(row['voice']).hexdigest())
    else:
        candidates=[x for x in detail['candidates'] if x.get('verified',True)]
        candidates.sort(key=lambda x:({'original':0,'mid':1,'thumb':2}.get(x['tier'],0),x['path']))
        selected=next((x for x in candidates if Path(x['path']).is_file()),None)
        if selected is None:
            return dict(**base,status='missing',reason='已关联的缓存文件现已不存在，请刷新媒体索引。')
        source=Path(selected['path']).resolve()
        if not source.is_relative_to(Path(account['path']).resolve()):
            raise ValueError('媒体路径超出当前账号目录。')
        if source.stat().st_size>200*1024*1024:
            return dict(**base,status='too_large',reason='单个媒体超过200MiB，未自动复制。')
        provenance=dict(source_path=str(source),source_bytes=source.stat().st_size,match=selected['match'],tier=selected['tier'])
        if row['kind']=='image':
            from PIL import Image
            reader=ImageReader(account)
            try:plain=reader.decode(source)
            finally:reader.close()
            if plain.startswith(b'wxgf'):
                import av
                helper=ImageReader(account)
                try:stream=helper.downloader._extract_hevc(plain)
                finally:helper.close()
                if not stream:raise ValueError('未解析出图片容器中的HEVC数据。')
                with av.open(io.BytesIO(stream),format='hevc') as c:
                    image=next(c.decode(video=0)).to_image()
                path=folder/'image.jpg';image.save(path,quality=95)
            else:
                with Image.open(io.BytesIO(plain)) as image:
                    image.load()
                    ext={'JPEG':'.jpg','PNG':'.png','WEBP':'.webp','GIF':'.gif','BMP':'.bmp'}.get(image.format)
                    if ext is None:raise ValueError('图片格式尚不支持。')
                    path=folder/('image'+ext);path.write_bytes(plain)
            with Image.open(path) as im:
                im.verify()
            with Image.open(path) as im:
                result=dict(status='ready',path=str(path),width=im.width,height=im.height,tier=selected['tier'],
                            original=selected['tier']=='original')
        else:
            suffix=source.suffix.lower()
            path=folder/('video.mp4' if row['kind']=='video' else 'attachment'+suffix)
            path=_copy_local(source,path)
            if row['kind']=='video':result=_video(path,folder,frames,transcribe,language)
            else:
                result=dict(status='ready',path=str(path),filename=source.name,document=_document(path,max_chars))
                (folder/'text.txt').write_text(result['document'].get('text',''),encoding='utf-8')
        with source.open('rb') as f:provenance['sha256']=hashlib.file_digest(f,'sha256').hexdigest()
    result={**base,**result,'provenance':provenance,'asset_signature':signature,'options':options}
    (folder/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    if 'transcription' in result:
        (folder/'transcript.txt').write_text(result['transcription']['text'],encoding='utf-8')
    with contextlib.closing(_db(write=True)) as c:
        c.execute('INSERT OR REPLACE INTO results VALUES(?,?)',(identity,json.dumps(result,ensure_ascii=False)));c.commit()
    return result


def search_media_content(keyword,session='',limit=20,offset=0):
    page(limit,offset)
    if not keyword.strip():
        raise ValueError('搜索词不能为空。')
    media_status()  # Verify the account fingerprint before joining results.
    # Apply the same live index and exclusion checks used by media retrieval.
    with contextlib.closing(connect()) as active, contextlib.closing(_db()) as c:
        c.execute('ATTACH DATABASE ? AS current_index',(INDEX.as_uri()+'?mode=ro',))
        where=['instr(lower(COALESCE(json_extract(r.result,\'$.transcription.text\'),\'\')||COALESCE(json_extract(r.result,\'$.document.text\'),\'\')),lower(?))>0']
        args=[keyword]
        if session:
            where.append('m.session_id=?');args.append(resolve_session(active,session)['username'])
        base=' FROM results r JOIN current_index.messages m ON m.identity=r.identity WHERE '+' AND '.join(where)
        rows=c.execute('SELECT r.identity,r.result'+base+' ORDER BY m.create_time DESC',args).fetchall()
    matches=[]
    for r in rows:
        try:_active(r['identity'])
        except ToolError:continue
        value=json.loads(r['result'])
        text=value.get('transcription',{}).get('text') or value.get('document',{}).get('text','')
        pos=text.lower().find(keyword.lower())
        matches.append(dict(identity=r['identity'],kind=value['kind'],session=value['session'],
            time=value['time'],excerpt=text[max(0,pos-100):pos+len(keyword)+300]))
    return dict(total=len(matches),offset=offset,matches=matches[offset:offset+limit],scope='已成功处理的语音、视频转写及附件正文')


if __name__=='__main__':
    p=argparse.ArgumentParser()
    p.add_argument('command',choices=['status','list','read'])
    p.add_argument('--session',default='');p.add_argument('--kind',default='')
    p.add_argument('--availability',default='');p.add_argument('--identity',default='')
    p.add_argument('--limit',type=int,default=20);p.add_argument('--no-transcribe',action='store_true')
    args=p.parse_args()
    result=media_status() if args.command=='status' else list_media(args.session,args.kind,args.availability,args.limit) if args.command=='list' else read_media(args.identity,transcribe=not args.no_transcribe)
    print(json.dumps(result,ensure_ascii=False,indent=2))
